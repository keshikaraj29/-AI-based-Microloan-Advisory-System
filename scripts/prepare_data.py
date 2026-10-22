#!/usr/bin/env python3
"""Reproducibly prepare local indexed evidence. Run from any directory.

  python scripts/prepare_data.py --loan-source /path/to/accepted_2007_to_2018Q4.csv.gz

The first run retains the original compressed loan archive under data/raw.
Later runs can omit --loan-source. Reads every valid loan, keeps a uniform
reservoir of 12,000 rows with seed 42, and imports all usable market rows.
No fitted model or workbook is loaded in web request handlers.
"""
import argparse
import csv
from contextlib import contextmanager,closing
from datetime import datetime, timezone
import gzip
import hashlib
import json
import math
import os
from pathlib import Path
import random
import shutil
import sqlite3
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault('LOKY_MAX_CPU_COUNT', str(min(os.cpu_count() or 1, 4)))
sys.path.insert(0, str(ROOT))
from research_analytics import build_summary, clean_narrative


def sha256(path):
    digest = hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def number(value):
    try:
        result = float(str(value).replace('%','').strip())
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def schema(conn):
    conn.executescript('''
      CREATE TABLE loans (
        id INTEGER PRIMARY KEY, source_row INTEGER NOT NULL,
        loan_amount REAL NOT NULL, term_months INTEGER, interest_rate REAL,
        annual_income REAL, dti REAL, purpose TEXT, state TEXT, issue_date TEXT,
        loan_status TEXT, description TEXT, cluster_id INTEGER,
        cluster_label TEXT, sentiment_score REAL, sentiment_label TEXT, topic_id INTEGER
      );
      CREATE TABLE market_prices (
        id INTEGER PRIMARY KEY, district TEXT, market TEXT, commodity TEXT,
        state TEXT, modal_price REAL NOT NULL, price_date TEXT NOT NULL,
        month INTEGER NOT NULL, district_key TEXT, market_key TEXT,
        commodity_key TEXT, state_key TEXT
      );
      CREATE TABLE farm_records (id INTEGER PRIMARY KEY, location TEXT, crop TEXT, season TEXT, row_json TEXT);
      CREATE TABLE dairy_records (id INTEGER PRIMARY KEY, location TEXT, product TEXT, row_json TEXT);
      CREATE TABLE locations (name TEXT PRIMARY KEY);
    ''')


def import_loans(conn, path, sample_size):
    rng = random.Random(42)
    reservoir = []
    valid, invalid, narratives, rows_seen = 0, 0, 0, 0
    with gzip.open(path,'rt',encoding='utf-8-sig',newline='') as f:
        reader = csv.DictReader(f)
        required = {'loan_amnt','term','int_rate','annual_inc','dti','purpose','loan_status','desc'}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError('Loan archive is missing required columns.')
        for index,row in enumerate(reader,1):
            rows_seen = index
            amount = number(row.get('loan_amnt'))
            if amount is None or amount <= 0:
                invalid += 1
                continue
            valid += 1
            if row.get('desc','').strip():
                narratives += 1
            # Algorithm R: every valid input record has sample_size/N chance.
            position = valid-1 if valid <= sample_size else rng.randrange(valid)
            if position < sample_size:
                term = number(row.get('term','').replace('months',''))
                income = number(row.get('annual_inc'))
                if income is not None and income < 0:
                    income = None
                record = (index, amount, int(term) if term else None,
                    number(row.get('int_rate')),income,number(row.get('dti')),
                    row.get('purpose','unknown').strip() or 'unknown',
                    row.get('addr_state','unknown').strip() or 'unknown',
                    row.get('issue_d',''),row.get('loan_status','Unknown'),
                    clean_narrative(row.get('desc','')))
                if valid <= sample_size:
                    reservoir.append(record)
                else:
                    reservoir[position] = record
            if index % 250000 == 0:
                print(f'Loan source scanned: {index:,} rows',flush=True)
    reservoir.sort(key=lambda row:row[0])
    conn.executemany('''INSERT INTO loans
       (id,source_row,loan_amount,term_months,interest_rate,annual_income,dti,purpose,state,issue_date,loan_status,description)
       VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''',[(i,*row) for i,row in enumerate(reservoir,1)])
    conn.execute('CREATE INDEX idx_loans_purpose ON loans(purpose,id)')
    conn.execute('CREATE INDEX idx_loans_state ON loans(state,id)')
    conn.execute('CREATE INDEX idx_loans_status ON loans(loan_status,id)')
    conn.commit()
    return {'rows_scanned':rows_seen,'valid_rows':valid,'invalid_loan_amount_rows':invalid,
      'sample_rows':len(reservoir),'nonempty_descriptions':narratives,
      'sampling':f'Uniform reservoir (Algorithm R), seed 42, target {sample_size:,}; complete file scanned once; eligibility: finite positive loan_amnt; no replacement; source rows sorted after selection.'}


def import_market(conn,path):
    import openpyxl
    book = openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    expected = ['District Name','Market Name','Commodity','Modal Price (Rs./Quintal)','Price Date','State']
    if any(field not in header for field in expected):
        raise ValueError('Market workbook headers differ from expected schema.')
    indices = {field:header.index(field) for field in expected}
    inserted, invalid, batch, locations = 0, 0, [], set()
    sql = '''INSERT INTO market_prices
      (district,market,commodity,state,modal_price,price_date,month,district_key,market_key,commodity_key,state_key)
      VALUES (?,?,?,?,?,?,?,?,?,?,?)'''
    for index,row in enumerate(rows,1):
        def val(field):
            return row[indices[field]]
        price = number(val('Modal Price (Rs./Quintal)'))
        date = val('Price Date')
        if isinstance(date,str):
            try:
                date = datetime.fromisoformat(date)
            except ValueError:
                try:
                    date = datetime.strptime(date,'%d/%m/%Y')
                except ValueError:
                    date = None
        if price is None or price <= 0 or not isinstance(date,datetime):
            invalid += 1
            continue
        district,market,commodity,state = [str(val(field) or '').strip()
            for field in ('District Name','Market Name','Commodity','State')]
        batch.append((district,market,commodity,state,price,date.date().isoformat(),date.month,
                      district.casefold(),market.casefold(),commodity.casefold(),state.casefold()))
        locations.update(x for x in (district,market,state) if x)
        inserted += 1
        if len(batch)>=5000:
            conn.executemany(sql,batch)
            batch.clear()
        if index % 100000 == 0:
            print(f'Market source imported: {index:,} rows',flush=True)
    if batch:
        conn.executemany(sql,batch)
    conn.executemany('INSERT OR IGNORE INTO locations(name) VALUES (?)',[(x,) for x in sorted(locations)])
    for field in ('district_key','market_key','state_key'):
        conn.execute(f'CREATE INDEX idx_market_{field} ON market_prices({field},commodity_key,month)')
    conn.execute('CREATE INDEX idx_market_commodity ON market_prices(commodity_key,id)')
    conn.commit()
    book.close()
    return {'rows_scanned':index,'imported_rows':inserted,'invalid_price_or_date_rows':invalid,
            'sheet':sheet.title,'price_unit':'INR per quintal','source_at_excel_row_limit':sheet.max_row==1048576}


def import_context_csv(conn,path,table):
    descriptions = 0
    with open(path,encoding='utf-8-sig',newline='') as f:
        rows = list(csv.DictReader(f))
    for i,row in enumerate(rows,1):
        descriptions += bool(row.get('Business Description','').strip())
        if table=='farm_records':
            conn.execute('INSERT INTO farm_records VALUES (?,?,?,?,?)',
                (i,row.get('Location'),row.get('Crops'),row.get('Season'),json.dumps(row,ensure_ascii=False)))
        else:
            conn.execute('INSERT INTO dairy_records VALUES (?,?,?,?)',
                (i,row.get('Location'),row.get('Product Name'),json.dumps(row,ensure_ascii=False)))
    if table=='farm_records':
        conn.execute('CREATE INDEX idx_farm_location_crop ON farm_records(location,crop,season)')
    else:
        conn.execute('CREATE INDEX idx_dairy_location ON dairy_records(location,product)')
    conn.commit()
    return {'rows':len(rows),'nonempty_business_descriptions':descriptions}


def install_research_database(staged_path,target_path):
    """Replace research tables transactionally while preserving local accounts.

    An existing database is backed up with SQLite's snapshot API, so WAL data
    is included. Stop the web server while preparing/recomputing research.
    Historical table updates never delete users, applications or payment logs.
    """
    staged_path,target_path = Path(staged_path),Path(target_path)
    if not target_path.exists():
        staged_path.replace(target_path)
        return None
    tables = ('loans','market_prices','farm_records','dairy_records','locations')
    backup_dir = target_path.parent/'backups'
    backup_dir.mkdir(exist_ok=True)
    backup_path = backup_dir/('pp1_before_rebuild_'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')+'.sqlite3')
    with closing(sqlite3.connect(target_path,timeout=30)) as conn:
        with closing(sqlite3.connect(backup_path)) as backup:
            conn.backup(backup)
        conn.execute('ATTACH DATABASE ? AS prepared',(str(staged_path.resolve()),))
        try:
            conn.execute('BEGIN IMMEDIATE')
            for table in tables:
                old_columns = [row[1] for row in conn.execute(f'PRAGMA main.table_info({table})')]
                new_columns = [row[1] for row in conn.execute(f'PRAGMA prepared.table_info({table})')]
                if old_columns!=new_columns:
                    raise RuntimeError(f'Research table {table} schema differs; existing database preserved. Backup: {backup_path.name}')
                conn.execute(f'DELETE FROM main.{table}')
                conn.execute(f'INSERT INTO main.{table} SELECT * FROM prepared.{table}')
            conn.execute('CREATE INDEX IF NOT EXISTS idx_market_commodity ON market_prices(commodity_key,id)')
            conn.execute('ANALYZE main')
            conn.commit()
        except BaseException:
            conn.rollback()
            raise
        finally:
            conn.execute('DETACH DATABASE prepared')
    staged_path.unlink()
    return str(backup_path)


@contextmanager
def preparation_lock(directory):
    """Prevent concurrent importers from deleting each other's staging file."""
    directory = Path(directory)
    directory.mkdir(parents=True,exist_ok=True)
    with open(directory/'.prepare_data.lock','a+b') as handle:
        if os.name=='nt':
            import msvcrt
            if handle.seek(0,2)==0:
                handle.write(b'0')
                handle.flush()
            handle.seek(0)
            try:
                msvcrt.locking(handle.fileno(),msvcrt.LK_NBLCK,1)
            except OSError as exc:
                raise RuntimeError('Another data preparation is already running. Wait for it to finish.') from exc
            try:
                yield
            finally:
                handle.seek(0)
                msvcrt.locking(handle.fileno(),msvcrt.LK_UNLCK,1)
        else:
            import fcntl
            try:
                fcntl.flock(handle.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB)
            except OSError as exc:
                raise RuntimeError('Another data preparation is already running. Wait for it to finish.') from exc
            try:
                yield
            finally:
                fcntl.flock(handle.fileno(),fcntl.LOCK_UN)


def _main_unlocked():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--loan-source',type=Path)
    parser.add_argument('--sample-size',type=int,default=12000)
    parser.add_argument('--analytics-only',action='store_true',help='Recompute analytics from an already imported database.')
    args = parser.parse_args()
    if args.sample_size < 100:
        parser.error('--sample-size must be at least 100.')
    data = ROOT / 'data'
    raw = data / 'raw'
    raw.mkdir(parents=True,exist_ok=True)
    db_path = data / 'pp1.sqlite3'
    manifest_path = data / 'source_manifest.json'
    if args.analytics_only:
        manifest = json.loads(manifest_path.read_text())
    else:
        archive = raw / 'accepted_2007_to_2018Q4.csv.gz'
        if args.loan_source and args.loan_source.resolve()!=archive.resolve():
            shutil.copy2(args.loan_source,archive)
        if not archive.is_file():
            parser.error('Pass --loan-source on the first run; compressed source is retained in data/raw.')
        temp_db = data / 'pp1.building.sqlite3'
        if temp_db.exists():
            temp_db.unlink()
        conn = sqlite3.connect(temp_db)
        conn.execute('PRAGMA journal_mode = DELETE')
        conn.execute('PRAGMA synchronous = NORMAL')
        schema(conn)
        started = time.perf_counter()
        loan_stats = import_loans(conn,archive,args.sample_size)
        market_path = ROOT / 'agmarknet_india_historical_prices_2024_2025.csv.xlsx'
        market_stats = import_market(conn,market_path)
        farm_stats = import_context_csv(conn,ROOT/'data_season.csv','farm_records')
        dairy_stats = import_context_csv(conn,ROOT/'dairy_dataset.csv','dairy_records')
        conn.execute('ANALYZE')
        check = conn.execute('PRAGMA integrity_check').fetchone()[0]
        if check!='ok':
            raise RuntimeError('Database integrity check failed: '+check)
        conn.close()
        backup_path = install_research_database(temp_db,db_path)
        sources = []
        specs = [(archive,loan_stats['valid_rows'],'USD','US LendingClub consumer lending; 2007–2018Q4'),
                 (market_path,market_stats['imported_rows'],'INR/quintal','Indian wholesale market prices; supplied 2024–2025 workbook'),
                 (ROOT/'data_season.csv',farm_stats['rows'],'INR where stated; mixed agriculture measures','Agricultural context; no linked loan outcomes'),
                 (ROOT/'dairy_dataset.csv',dairy_stats['rows'],'INR','Dairy context; no linked loan outcomes')]
        for path,count,currency,scope in specs:
            sources.append({'name':path.name,'path':str(path.relative_to(ROOT)),
                'records':count,'bytes':path.stat().st_size,'sha256':sha256(path),
                'currency':currency,'scope':scope,'origin':'User-supplied local file; upstream license/collection not independently verified.'})
        manifest = {'schema_version':1,'generated_at':datetime.now(timezone.utc).isoformat(),
            'loans':loan_stats,'market':market_stats,'farm':farm_stats,'dairy':dairy_stats,
            'sources':sources,'build_seconds':round(time.perf_counter()-started,3),'integrity_check':check,
            'pre_rebuild_backup':str(Path(backup_path).relative_to(ROOT)) if backup_path else None,
            'privacy':'Web rows use local sample IDs and omit original loan IDs, URLs, job titles and raw descriptions.'}
        manifest_path.write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print('Computing offline cluster, narrative and outcome analytics...',flush=True)
    summary = build_summary(db_path,manifest)
    path = data / 'research_summary.json'
    temporary = path.with_suffix('.building.json')
    temporary.write_text(json.dumps(summary,indent=2,ensure_ascii=False,allow_nan=False)+'\n')
    temporary.replace(path)
    print(json.dumps({'database':str(db_path),'summary':str(path),'dataset':summary['dataset'],
                      'metrics':summary['metrics']},indent=2),flush=True)


def main():
    with preparation_lock(ROOT/'data'):
        _main_unlocked()


if __name__=='__main__':
    main()
